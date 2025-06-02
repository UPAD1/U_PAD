import os
import fitz  # PyMuPDF
import uuid
import tempfile
import subprocess
import numpy as np
import cv2
from pptx import Presentation
import docx

from app.utils.vision_ocr import extract_text_from_image_gvision


def is_scanned_pdf(page):
    return not page.get_text().strip()


def extract_text_from_docx(file_path):
    try:
        doc = docx.Document(file_path)
        text = "\n".join(p.text for p in doc.paragraphs)
        return text if text.strip() else "[⚠️ DOCX에서 텍스트를 추출하지 못했습니다.]"
    except Exception as e:
        return f"[❌ DOCX 처리 오류: {e}]"


def extract_text_from_pptx(file_path):
    try:
        prs = Presentation(file_path)
        text_runs = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    text_runs.append(shape.text)
        return "\n".join(text_runs)
    except Exception as e:
        return f"[❌ PPTX 처리 오류: {e}]"


def extract_text_from_hwp(file_path):
    try:
        with tempfile.NamedTemporaryFile(delete=False, suffix=".txt") as tmp:
            subprocess.run(['hwp5txt', file_path], stdout=tmp, stderr=subprocess.DEVNULL)
        with open(tmp.name, 'r', encoding='cp949') as f:
            return f.read()
    except Exception as e:
        return f"[❌ HWP 처리 오류: {e}]"


def process_document(file_path: str):
    output_dir = "static/uploads/document"
    os.makedirs(output_dir, exist_ok=True)
    result_image_paths = []
    output_text = ""
    ext = os.path.splitext(file_path)[1].lower()

    try:
        if ext == '.pdf':
            doc = fitz.open(file_path)
            for i in range(len(doc)):
                page = doc.load_page(i)
                if is_scanned_pdf(page):
                    pix = page.get_pixmap(dpi=300)
                    img_path = os.path.join(output_dir, f"{uuid.uuid4().hex}_page{i}.jpg")
                    pix.save(img_path)
                    result_image_paths.append(img_path)

                    extracted_text = extract_text_from_image_gvision(img_path)
                else:
                    extracted_text = page.get_text()

                output_text += f"\n\n--- Page {i+1} ---\n{extracted_text.strip() or '[⚠️ 텍스트 없음]'}"

        elif ext == '.docx':
            output_text = extract_text_from_docx(file_path)
        elif ext == '.pptx':
            output_text = extract_text_from_pptx(file_path)
        elif ext == '.hwp':
            output_text = extract_text_from_hwp(file_path)
        else:
            output_text = "⚠️ 지원하지 않는 파일 형식입니다."

    except Exception as e:
        output_text = f"[❌ 문서 전체 처리 실패: {e}]"

    return result_image_paths, output_text


# (필요 시) 이미지 객체에서도 텍스트 추출용
def extract_text_from_image(image: np.ndarray) -> str:
    try:
        # 프레임을 임시 이미지로 저장 후 처리
        with tempfile.NamedTemporaryFile(suffix=".jpg", delete=True) as tmp_img:
            cv2.imwrite(tmp_img.name, image)
            return extract_text_from_image_gvision(tmp_img.name)
    except Exception as e:
        return f"[❌ 이미지 OCR 실패: {e}]"
